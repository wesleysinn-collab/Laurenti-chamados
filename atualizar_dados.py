#!/usr/bin/env python3
"""
Gera o arquivo dados.js a partir das planilhas de lojas e de selfs.

Uso:
    pip install openpyxl
    python atualizar_dados.py "Cadastro de Lojas.xlsx" "Tabela selfs.xlsx"

Depois é só dar commit no dados.js (o index.html lê esse arquivo).
"""
import re, sys, json, collections
import openpyxl

def limpa(v):
    if v is None: return ''
    return re.sub(r'\s+', ' ', str(v).replace('\t', ' ')).strip()

def cep_fmt(v):
    d = re.sub(r'\D', '', limpa(v))
    return f'{d[:5]}-{d[5:]}' if len(d) == 8 else limpa(v)

def numero(v):
    s = limpa(v)
    return s[:-2] if s.endswith('.0') else s

def num_loja(txt):
    m = re.search(r'(\d+)', str(txt)); return int(m.group(1)) if m else None

def bal(v):
    s = limpa(v)
    if re.fullmatch(r'\d+\s*[xX]\s*\d+', s): return s.lower().replace(' ', '')
    if s in ('', '-', '--', '- -', '*'): return ''
    return s

# Erros de digitação conhecidos nas planilhas
CORRIGE_CIDADE = {'Barra Vellha': 'Barra Velha', 'Balneario Piçarras': 'Balneário Piçarras',
                  'Jaragua do Sul': 'Jaraguá do Sul', 'Balneario Camboriu': 'Balneário Camboriú'}

def le_lojas(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    # Cidade: a coluna CIDADE da aba SUPERMERCADO vem com #VALUE!, então usamos a aba de localidades
    loc = {}
    for r in wb['LOJAS - APENAS LOCALIDADE'].iter_rows(min_row=2, values_only=True):
        if isinstance(r[0], int): loc[r[0]] = limpa(r[3])
        if isinstance(r[5], int): loc.setdefault(r[5], limpa(r[8]))
    lojas = {}
    for r in wb['SUPERMERCADO'].iter_rows(min_row=2, values_only=True):
        n = r[1]
        if not isinstance(n, int): continue
        cid = limpa(r[8])
        origem = 'planilha'
        if cid.startswith('#') or not cid:
            cid = loc.get(n, ''); origem = 'planilha'
        lojas[n] = dict(n=n, tipo=limpa(r[2]), logradouro=limpa(r[5]), numero=numero(r[6]),
                        bairro=limpa(r[7]), cidade=cid, uf=limpa(r[9]) or 'SC',
                        cep=cep_fmt(r[10]), sala='', origem=origem)
    for r in wb['FARMÁCIA'].iter_rows(min_row=3, values_only=True):
        n = r[0]
        if not isinstance(n, int): continue
        lojas[n] = dict(n=n, tipo=limpa(r[1]), logradouro=limpa(r[5]), numero=numero(r[6]),
                        bairro=limpa(r[8]), cidade=limpa(r[9]), uf=limpa(r[10]) or 'SC',
                        cep=cep_fmt(r[11]), sala=numero(r[7]), ref=r[4], origem='planilha')
    # Cidade que ainda ficou vazia: se o CEP é genérico da cidade (termina em -000) e
    # outras lojas com o MESMO CEP concordam, usa essa cidade (marcada como "inferida").
    por_cep = collections.defaultdict(set)
    for l in lojas.values():
        if l['cidade'] and l['cep']: por_cep[l['cep']].add(l['cidade'])
    for l in lojas.values():
        if not l['cidade'] and l['cep'].endswith('-000') and len(por_cep[l['cep']]) == 1:
            l['cidade'] = next(iter(por_cep[l['cep']])); l['origem'] = 'inferida'
    for l in lojas.values():
        l['cidade'] = CORRIGE_CIDADE.get(l['cidade'], l['cidade'])
    return [lojas[k] for k in sorted(lojas)]

def le_selfs(caminho):
    ws = openpyxl.load_workbook(caminho, data_only=True)['Selfs']
    atual = None; out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is not None: atual = num_loja(r[0])
        if not isinstance(r[1], int) or r[2] not in ('Self', 'Flex'):
            continue  # ignora cancelas, PC apoio etc.
        serie = limpa(r[11])
        if not serie or set(serie) <= set('Xx?'): serie = ''
        out.append(dict(loja=atual, pdv=r[1], tipo=limpa(r[2]), serie=serie, bal=bal(r[3])))
    return out

if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    lojas, selfs = le_lojas(sys.argv[1]), le_selfs(sys.argv[2])
    with open('dados.js', 'w', encoding='utf-8') as f:
        f.write('// Gerado por atualizar_dados.py — não editar à mão\n')
        f.write('const LOJAS = [\n' + ',\n'.join(json.dumps(l, ensure_ascii=False) for l in lojas) + '\n];\n')
        f.write('const SELFS = [\n' + ',\n'.join(json.dumps(x, ensure_ascii=False) for x in selfs) + '\n];\n')
    sem_cid = [l['n'] for l in lojas if not l['cidade']]
    import os
    if not os.path.exists('correcoes.js'):   # nunca sobrescreve correções já feitas
        with open('correcoes.js', 'w', encoding='utf-8') as f:
            f.write('// Correções manuais que valem POR CIMA da planilha. Edite aqui direto no GitHub.\n')
            f.write('// Exemplo:  34: { cidade: "Timbó" },\n')
            f.write('const CORRECOES = {\n')
            for l in lojas:
                if not l['cidade']:
                    f.write(f"  // Loja {l['n']} ({l['tipo']}) - {l['logradouro']}, {l['numero']} - {l['bairro']} - CEP {l['cep']}\n")
                    f.write(f"  {l['n']}: {{ cidade: \"\" }},\n")
            f.write('};\n')
    print(f'{len(lojas)} lojas, {len(selfs)} selfs/flex gravados em dados.js')
    if sem_cid: print('Lojas SEM cidade (preencher na planilha de localidades):', sem_cid)
